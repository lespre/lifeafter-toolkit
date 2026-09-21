# -*- coding: utf-8 -*-
"""读取docx文件内容"""
import docx
import sys

docx_path = r'E:\提取成果\lifeafter拆包思路与进展 .docx'

try:
    doc = docx.Document(docx_path)
    
    print(f"=== 文档标题 ===")
    print(f"段落数: {len(doc.paragraphs)}")
    print(f"表格数: {len(doc.tables)}")
    print()
    
    print("=== 文档内容 ===")
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            # 获取段落样式
            style = para.style.name if para.style else "Normal"
            if style.startswith("Heading"):
                print(f"\n{'='*60}")
                print(f"[{style}] {para.text}")
                print(f"{'='*60}")
            else:
                print(para.text)
    
    print("\n\n=== 表格内容 ===")
    for table_idx, table in enumerate(doc.tables):
        print(f"\n--- 表格 {table_idx+1} ({len(table.rows)}行 x {len(table.columns)}列) ---")
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            print(f"  行{row_idx}: {' | '.join(cells)}")
            
except Exception as e:
    print(f"读取失败: {e}")
    import traceback
    traceback.print_exc()
